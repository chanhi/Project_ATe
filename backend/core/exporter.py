"""
테스트 케이스 내보내기 (7가지 포맷)
───────────────────────────────────
json, csv, xlsx, md, yaml, playwright, html
"""

import csv
import io
import json
from typing import Any


def export_to_json(test_cases: list[dict]) -> str:
    """JSON 포맷 내보내기"""
    clean = []
    for tc in test_cases:
        clean.append({
            "test_case_id": tc["test_case_id"],
            "title": tc["title"],
            "description": tc.get("description", ""),
            "precondition": tc.get("precondition", ""),
            "steps": tc.get("steps", []),
            "expected_result": tc.get("expected_result", ""),
            "priority": tc.get("priority", "medium"),
            "category": tc.get("category", ""),
            "technique": tc.get("technique", ""),
            "target_urls": tc.get("target_urls", []),
        })
    return json.dumps(clean, ensure_ascii=False, indent=2)


def export_to_csv(test_cases: list[dict]) -> str:
    """CSV 포맷 — 스프레드시트 호환"""
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([
        "테스트 ID", "제목", "설명", "사전조건",
        "스텝", "기대결과", "우선순위", "카테고리", "기법",
    ])

    for tc in test_cases:
        steps_str = " → ".join(
            f"[{s.get('step_no', i+1)}] {s.get('action', '')} {s.get('target', '')} {s.get('input', '')}"
            for i, s in enumerate(tc.get("steps", []))
        )
        writer.writerow([
            tc["test_case_id"],
            tc["title"],
            tc.get("description", ""),
            tc.get("precondition", ""),
            steps_str,
            tc.get("expected_result", ""),
            tc.get("priority", "medium"),
            tc.get("category", ""),
            tc.get("technique", ""),
        ])

    return output.getvalue()


def export_to_xlsx(test_cases: list[dict]) -> bytes:
    """Excel — 스텝별 행 분리 + 헤더 스타일"""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        raise RuntimeError("openpyxl 미설치. pip install openpyxl")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "테스트 케이스"

    headers = [
        "테스트 ID", "제목", "설명", "사전조건",
        "스텝번호", "액션", "대상", "입력값", "기대결과",
        "우선순위", "카테고리", "기법",
    ]
    ws.append(headers)

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="2E7D32", end_color="2E7D32", fill_type="solid")

    for tc in test_cases:
        steps = tc.get("steps") or [{}]
        for i, step in enumerate(steps):
            ws.append([
                tc["test_case_id"] if i == 0 else "",
                tc["title"] if i == 0 else "",
                tc.get("description", "") if i == 0 else "",
                tc.get("precondition", "") if i == 0 else "",
                step.get("step_no", i + 1),
                step.get("action", ""),
                step.get("target", ""),
                step.get("input", ""),
                step.get("expected", tc.get("expected_result", "")),
                tc.get("priority", "") if i == 0 else "",
                tc.get("category", "") if i == 0 else "",
                tc.get("technique", "") if i == 0 else "",
            ])

    # 열 너비 자동
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def export_to_markdown(test_cases: list[dict]) -> str:
    """Markdown — 문서화/리뷰용"""
    lines = ["# 테스트 케이스\n"]

    for tc in test_cases:
        lines.append(f"## {tc['test_case_id']}: {tc['title']}\n")

        if tc.get("description"):
            lines.append(f"{tc['description']}\n")

        lines.append(f"- **우선순위**: {tc.get('priority', 'medium')}")
        lines.append(f"- **기법**: {tc.get('technique', '-')}")
        lines.append(f"- **카테고리**: {tc.get('category', '-')}")

        if tc.get("precondition"):
            lines.append(f"- **사전조건**: {tc['precondition']}")

        lines.append("\n### 테스트 스텝\n")
        lines.append("| 번호 | 액션 | 대상 | 입력값 | 기대결과 |")
        lines.append("|------|------|------|--------|----------|")

        for i, step in enumerate(tc.get("steps", [])):
            lines.append(
                f"| {step.get('step_no', i+1)} "
                f"| {step.get('action', '')} "
                f"| {step.get('target', '')} "
                f"| {step.get('input', '')} "
                f"| {step.get('expected', '')} |"
            )

        if tc.get("expected_result"):
            lines.append(f"\n**최종 기대결과**: {tc['expected_result']}")

        lines.append("\n---\n")

    return "\n".join(lines)


def export_to_yaml(test_cases: list[dict]) -> str:
    """YAML 포맷"""
    try:
        import yaml
    except ImportError:
        return _yaml_fallback(test_cases)

    clean = []
    for tc in test_cases:
        clean.append({
            "test_case_id": tc["test_case_id"],
            "title": tc["title"],
            "description": tc.get("description", ""),
            "precondition": tc.get("precondition", ""),
            "steps": tc.get("steps", []),
            "expected_result": tc.get("expected_result", ""),
            "priority": tc.get("priority", "medium"),
            "technique": tc.get("technique", ""),
        })

    return yaml.dump(clean, allow_unicode=True, default_flow_style=False, sort_keys=False)


def _yaml_fallback(test_cases: list[dict]) -> str:
    """PyYAML 없을 때 최소 YAML 생성"""
    lines = []
    for tc in test_cases:
        lines.append(f"- test_case_id: \"{tc['test_case_id']}\"")
        lines.append(f"  title: \"{tc['title']}\"")
        lines.append(f"  priority: \"{tc.get('priority', 'medium')}\"")
        lines.append(f"  technique: \"{tc.get('technique', '')}\"")
        lines.append(f"  steps:")
        for s in tc.get("steps", []):
            lines.append(f"    - action: \"{s.get('action', '')}\"")
            lines.append(f"      target: \"{s.get('target', '')}\"")
            lines.append(f"      input: \"{s.get('input', '')}\"")
        lines.append("")
    return "\n".join(lines)


def export_to_playwright(test_cases: list[dict], base_url: str = "") -> str:
    """Playwright TypeScript 코드 내보내기"""
    lines = ["import { test, expect } from '@playwright/test';", ""]

    for tc in test_cases:
        title = tc["title"].replace("'", "\\'")
        lines.append(f"test('{title}', async ({{ page }}) => {{")

        if base_url:
            lines.append(f"  await page.goto('{base_url}');")

        for step in tc.get("steps", []):
            action = step.get("action", "").lower()
            target = step.get("target", "")
            value = step.get("input", "")

            if action in ("fill", "입력"):
                lines.append(f"  await page.fill('{target}', '{value}');")
            elif action in ("click", "클릭"):
                lines.append(f"  await page.click('{target}');")
            elif action in ("goto", "이동"):
                lines.append(f"  await page.goto('{target or value}');")
            elif action in ("assert_text", "텍스트확인"):
                lines.append(f"  await expect(page.locator('{target}')).toContainText('{value}');")
            elif action in ("assert_url", "url확인"):
                lines.append(f"  await expect(page).toHaveURL('{value}');")
            elif action in ("assert_visible", "표시확인"):
                lines.append(f"  await expect(page.locator('{target}')).toBeVisible();")
            elif action in ("wait", "대기"):
                lines.append(f"  await page.waitForTimeout({value or 1000});")
            elif action in ("screenshot", "스크린샷"):
                lines.append(f"  await page.screenshot({{ path: '{value or 'screenshot.png'}' }});")
            else:
                lines.append(f"  // {action}: {target} {value}")

        lines.append("});")
        lines.append("")

    return "\n".join(lines)


def export_to_html(test_cases: list[dict]) -> str:
    """HTML 보고서 — 그대로 브라우저에서 열람 가능"""
    rows = []
    for tc in test_cases:
        steps_html = "<ol>"
        for s in tc.get("steps", []):
            steps_html += f"<li><b>{s.get('action','')}</b> → {s.get('target','')} ({s.get('input','')})</li>"
        steps_html += "</ol>"

        priority_color = {
            "critical": "#d32f2f", "high": "#f57c00",
            "medium": "#1976d2", "low": "#388e3c",
        }.get(tc.get("priority", "medium"), "#666")

        rows.append(f"""
        <tr>
          <td>{tc['test_case_id']}</td>
          <td><b>{tc['title']}</b><br><small>{tc.get('description','')}</small></td>
          <td>{steps_html}</td>
          <td>{tc.get('expected_result','')}</td>
          <td style="color:{priority_color}; font-weight:bold">{tc.get('priority','medium')}</td>
          <td>{tc.get('technique','')}</td>
        </tr>""")

    return f"""<!DOCTYPE html>
<html lang="ko">
<head>
<meta charset="UTF-8">
<title>ATe 테스트 케이스 보고서</title>
<style>
  body {{ font-family: system-ui, 'Pretendard', sans-serif; margin: 40px; background: #f5f5f5; }}
  h1 {{ color: #1a1a2e; }}
  table {{ width: 100%; border-collapse: collapse; background: white; box-shadow: 0 2px 8px rgba(0,0,0,0.1); }}
  th {{ background: #00D47B; color: white; padding: 12px; text-align: left; }}
  td {{ padding: 10px 12px; border-bottom: 1px solid #eee; vertical-align: top; }}
  tr:hover {{ background: #f0fff4; }}
  ol {{ margin: 4px 0; padding-left: 20px; }}
  small {{ color: #888; }}
</style>
</head>
<body>
<h1>ATe 테스트 케이스 보고서</h1>
<p>총 <b>{len(test_cases)}</b>개 케이스 | 생성일: <span id="date"></span></p>
<table>
<thead><tr><th>ID</th><th>제목</th><th>스텝</th><th>기대결과</th><th>우선순위</th><th>기법</th></tr></thead>
<tbody>{''.join(rows)}</tbody>
</table>
<script>document.getElementById('date').textContent=new Date().toLocaleDateString('ko-KR');</script>
</body>
</html>"""


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
#  통합 Export 디스패처
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

EXPORTERS = {
    "json":       (export_to_json,       "application/json",  ".json"),
    "csv":        (export_to_csv,        "text/csv",          ".csv"),
    "xlsx":       (export_to_xlsx,
                   "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                   ".xlsx"),
    "md":         (export_to_markdown,   "text/markdown",     ".md"),
    "yaml":       (export_to_yaml,       "text/yaml",         ".yaml"),
    "playwright": (export_to_playwright, "text/plain",        ".spec.ts"),
    "html":       (export_to_html,       "text/html",         ".html"),
}


def export_test_cases(test_cases: list[dict], fmt: str, **kwargs) -> tuple[Any, str, str]:
    """
    Returns: (content, content_type, file_extension)
    content는 str 또는 bytes
    """
    if fmt not in EXPORTERS:
        raise ValueError(f"지원하지 않는 포맷: {fmt}. 지원: {list(EXPORTERS.keys())}")

    func, content_type, ext = EXPORTERS[fmt]
    if fmt == "playwright":
        content = func(test_cases, **kwargs)
    else:
        content = func(test_cases)
    return content, content_type, ext
