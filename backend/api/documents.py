"""
문서(기획서) 업로드 API
───────────────────────
업로드 시 자동으로 텍스트 추출 + DB에 저장.

POST   /documents/upload         → 파일 업로드 + 자동 텍스트 추출
PUT    /documents/{id}/extract   → 수동으로 텍스트 업데이트 (선택)
GET    /documents/{id}           → 문서 메타/텍스트 조회
GET    /documents                → 목록 조회
DELETE /documents/{id}           → 삭제
"""

import os
import io
import uuid
import logging
from datetime import datetime, timezone

from fastapi import APIRouter, HTTPException, Query, UploadFile, File, Form

from core.config import get_settings
from core.mongo import (
    get_projects_collection,
    get_documents_collection,
    get_test_cases_collection,
)
from schemas.schemas import (
    APIResponse,
    UploadedDocumentDoc,
    DocumentExtractUpdateRequest,
)

logger = logging.getLogger(__name__)
settings = get_settings()
router = APIRouter(prefix="/documents", tags=["Documents"])

ALLOWED_EXTENSIONS = {"pdf", "docx", "xlsx", "txt", "md", "hwp", "pptx", "csv"}


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
#  텍스트 추출 유틸
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

def _extract_text_from_file(content: bytes, file_ext: str, filename: str = "") -> tuple[str, str | None]:
    """
    파일 바이트에서 텍스트를 추출.
    반환: (extracted_text, error_message)
    """
    try:
        # 텍스트 파일 (txt, md, csv)
        if file_ext in ("txt", "md", "csv"):
            for encoding in ("utf-8", "cp949", "euc-kr", "latin-1"):
                try:
                    return content.decode(encoding), None
                except UnicodeDecodeError:
                    continue
            return content.decode("utf-8", errors="replace"), None

        # PDF
        if file_ext == "pdf":
            try:
                import fitz  # PyMuPDF
                text = ""
                with fitz.open(stream=content, filetype="pdf") as doc:
                    for page in doc:
                        text += page.get_text()
                return text, None
            except ImportError:
                return "", "PyMuPDF(fitz)가 설치되지 않음"
            except Exception as e:
                return "", f"PDF 파싱 실패: {e}"

        # DOCX
        if file_ext == "docx":
            try:
                from docx import Document
                doc = Document(io.BytesIO(content))
                text = "\n".join(p.text for p in doc.paragraphs)
                # 테이블 텍스트도 추가
                for table in doc.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            if cell.text.strip():
                                text += "\n" + cell.text
                return text, None
            except ImportError:
                return "", "python-docx가 설치되지 않음 (pip install python-docx)"
            except Exception as e:
                return "", f"DOCX 파싱 실패: {e}"

        # XLSX
        if file_ext == "xlsx":
            try:
                from openpyxl import load_workbook
                wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
                text = ""
                for sheet in wb.worksheets:
                    text += f"\n=== Sheet: {sheet.title} ===\n"
                    for row in sheet.iter_rows(values_only=True):
                        row_text = " | ".join(str(c) if c is not None else "" for c in row)
                        if row_text.strip(" |"):
                            text += row_text + "\n"
                return text, None
            except ImportError:
                return "", "openpyxl이 설치되지 않음"
            except Exception as e:
                return "", f"XLSX 파싱 실패: {e}"

        # PPTX
        if file_ext == "pptx":
            try:
                from pptx import Presentation
                prs = Presentation(io.BytesIO(content))
                text = ""
                for i, slide in enumerate(prs.slides, 1):
                    text += f"\n=== Slide {i} ===\n"
                    for shape in slide.shapes:
                        if hasattr(shape, "text") and shape.text:
                            text += shape.text + "\n"
                return text, None
            except ImportError:
                return "", "python-pptx가 설치되지 않음"
            except Exception as e:
                return "", f"PPTX 파싱 실패: {e}"

        # HWP (한글 문서) - 라이브러리 의존성이 까다로워서 일단 미지원
        if file_ext == "hwp":
            return "", "HWP 파일은 현재 자동 추출 미지원. txt로 변환해서 업로드해주세요."

        return "", f"지원하지 않는 형식: {file_ext}"

    except Exception as e:
        logger.exception("텍스트 추출 중 예외")
        return "", f"추출 실패: {e}"


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
#  업로드 (자동 텍스트 추출 포함)
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

@router.post(
    "/upload",
    response_model=APIResponse,
    status_code=201,
    summary="기획서/문서 업로드 (자동 텍스트 추출)",
)
async def upload_document(
    project_id: str = Form(...),
    file: UploadFile = File(...),
):
    """
    파일 업로드 시 자동으로 텍스트를 추출하여 DB에 저장한다.

    지원 자동 추출:
    - PDF, DOCX, XLSX, PPTX, TXT, MD, CSV

    추출 실패 시 메타데이터만 저장되고 extracted_text는 비어있다.
    """
    projects = get_projects_collection()
    documents = get_documents_collection()

    # 프로젝트 확인
    project = await projects.find_one({"project_id": project_id})
    if not project:
        raise HTTPException(
            status_code=404,
            detail={
                "status": "error",
                "message": f"프로젝트를 찾을 수 없습니다: {project_id}",
                "error_code": "PROJECT_NOT_FOUND",
            },
        )

    # 확장자 검증
    ext = file.filename.rsplit(".", 1)[-1].lower() if "." in (file.filename or "") else ""
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail={
                "status": "error",
                "message": f"지원하지 않는 파일 형식: .{ext}",
                "error_code": "UNSUPPORTED_FILE_TYPE",
            },
        )

    # 파일 읽기 + 크기 검증
    content = await file.read()
    if len(content) > settings.MAX_UPLOAD_SIZE_MB * 1024 * 1024:
        raise HTTPException(
            status_code=400,
            detail={
                "status": "error",
                "message": f"파일 크기가 {settings.MAX_UPLOAD_SIZE_MB}MB를 초과합니다.",
                "error_code": "FILE_TOO_LARGE",
            },
        )

    # 파일 디스크 저장
    document_id = f"doc-{uuid.uuid4().hex[:8]}"
    os.makedirs(settings.UPLOAD_DIR, exist_ok=True)
    safe_name = "".join(c for c in file.filename if c.isalnum() or c in "._-가-힣 ").strip()
    file_path = os.path.join(settings.UPLOAD_DIR, f"{document_id}_{safe_name}")
    with open(file_path, "wb") as f:
        f.write(content)

    # ── 자동 텍스트 추출 ──
    extracted_text, extract_error = _extract_text_from_file(content, ext, file.filename)
    extract_status = "extracted" if extracted_text else ("failed" if extract_error else "empty")

    if extract_error:
        logger.warning(f"문서 {document_id} 텍스트 추출 실패: {extract_error}")

    # DB 저장
    doc = UploadedDocumentDoc(
        document_id=document_id,
        project_id=project_id,
        filename=file.filename,
        file_type=ext,
        file_path=file_path,
        file_size_bytes=len(content),
        extracted_text=extracted_text or None,
    )
    await documents.insert_one(doc.model_dump())

    response_data = {
        "document_id": document_id,
        "filename": file.filename,
        "file_type": ext,
        "file_size_bytes": len(content),
        "extract_status": extract_status,
        "extracted_length": len(extracted_text),
    }
    if extract_error:
        response_data["extract_error"] = extract_error
    if extracted_text:
        response_data["preview"] = extracted_text[:300] + ("..." if len(extracted_text) > 300 else "")

    if extracted_text:
        message = f"업로드 및 텍스트 추출 완료 ({len(extracted_text):,}자). 이제 시나리오 생성이 가능합니다."
    elif extract_error:
        message = f"업로드는 완료됐지만 텍스트 추출 실패: {extract_error}"
    else:
        message = "업로드 완료. 텍스트가 비어있어 자연어 입력으로 시나리오를 생성하세요."

    return APIResponse(
        status="success",
        data=response_data,
        message=message,
    )


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
#  수동 텍스트 업데이트 (백업용)
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

@router.put(
    "/{document_id}/extract",
    response_model=APIResponse,
    summary="텍스트 수동 업데이트",
)
async def update_extracted_text(
    document_id: str,
    request: DocumentExtractUpdateRequest,
):
    """자동 추출이 실패했거나 텍스트를 직접 입력하고 싶을 때 사용."""
    documents = get_documents_collection()

    result = await documents.update_one(
        {"document_id": document_id},
        {"$set": {"extracted_text": request.extracted_text}},
    )

    if result.matched_count == 0:
        raise HTTPException(
            status_code=404,
            detail=f"문서를 찾을 수 없습니다: {document_id}",
        )

    return APIResponse(
        status="success",
        data={
            "document_id": document_id,
            "extracted_length": len(request.extracted_text),
        },
        message="extracted_text 업데이트 완료",
    )


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
#  텍스트 재추출 (자동 추출이 실패했을 때 다시 시도)
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

@router.post(
    "/{document_id}/re-extract",
    response_model=APIResponse,
    summary="저장된 파일에서 텍스트 재추출",
)
async def re_extract_text(document_id: str):
    """디스크에 저장된 원본 파일을 다시 읽어 텍스트를 추출한다."""
    documents = get_documents_collection()

    doc = await documents.find_one({"document_id": document_id})
    if not doc:
        raise HTTPException(404, "문서를 찾을 수 없습니다.")

    file_path = doc.get("file_path")
    if not file_path or not os.path.exists(file_path):
        raise HTTPException(400, "원본 파일을 찾을 수 없습니다.")

    with open(file_path, "rb") as f:
        content = f.read()

    extracted_text, error = _extract_text_from_file(
        content, doc.get("file_type", ""), doc.get("filename", "")
    )

    if not extracted_text:
        raise HTTPException(500, f"재추출 실패: {error}")

    await documents.update_one(
        {"document_id": document_id},
        {"$set": {"extracted_text": extracted_text}},
    )

    return APIResponse(
        status="success",
        data={
            "document_id": document_id,
            "extracted_length": len(extracted_text),
            "preview": extracted_text[:300],
        },
        message=f"재추출 완료 ({len(extracted_text):,}자)",
    )


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
#  조회 / 삭제
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

@router.get("/{document_id}", response_model=APIResponse)
async def get_document(document_id: str):
    documents = get_documents_collection()
    doc = await documents.find_one({"document_id": document_id}, {"_id": 0})
    if not doc:
        raise HTTPException(404, "문서를 찾을 수 없습니다.")
    if doc.get("created_at"):
        doc["created_at"] = doc["created_at"].isoformat()
    return APIResponse(status="success", data=doc)


@router.get("", response_model=APIResponse)
async def list_documents(
    project_id: str | None = Query(None),
    limit: int = Query(50, ge=1, le=200),
    offset: int = Query(0, ge=0),
):
    documents = get_documents_collection()
    query = {}
    if project_id:
        query["project_id"] = project_id

    cursor = documents.find(query, {"_id": 0, "extracted_text": 0}).sort("created_at", -1).skip(offset).limit(limit)
    docs = await cursor.to_list(length=limit)

    for doc in docs:
        if doc.get("created_at"):
            doc["created_at"] = doc["created_at"].isoformat()

    total = await documents.count_documents(query)

    return APIResponse(
        status="success",
        data={"items": docs, "total": total, "limit": limit, "offset": offset},
    )


@router.delete("/{document_id}", response_model=APIResponse)
async def delete_document(document_id: str):
    documents = get_documents_collection()
    test_cases = get_test_cases_collection()

    doc = await documents.find_one({"document_id": document_id})
    if not doc:
        raise HTTPException(404, "문서를 찾을 수 없습니다.")

    file_path = doc.get("file_path")
    if file_path and os.path.exists(file_path):
        try:
            os.remove(file_path)
        except OSError:
            pass

    await documents.delete_one({"document_id": document_id})
    await test_cases.update_many(
        {"document_id": document_id},
        {"$set": {"document_id": None}},
    )

    return APIResponse(
        status="success",
        data={"document_id": document_id},
        message="문서 삭제 완료",
    )